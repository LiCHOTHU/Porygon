#!/usr/bin/env python3
"""Build editable CAST video slides and a PDF from the same layout primitives."""
from __future__ import annotations
import argparse
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import subprocess
import sys
import textwrap

# Optional temporary dependency directory; normal virtualenv installs also work.
DEPS = Path(os.environ.get('ICRA_VIDEO_DEPS', '/tmp/icra-video-deps'))
if DEPS.exists():
    sys.path.insert(0, str(DEPS))
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Pt
from pptx.oxml.xmlchemy import OxmlElement
from reportlab.pdfgen import canvas as pdfcanvas
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
import matplotlib
matplotlib.use('Agg')
from matplotlib import pyplot as plt
from matplotlib.font_manager import findfont, FontProperties
import numpy as np

ROOT = Path(__file__).resolve().parents[1]
W, H = 960, 540
BG, CARD, PANEL = '#0D1827', '#16273B', '#102135'
WHITE, MUTED, LINE = '#F3F6FA', '#AFBDD0', '#33455B'
TEAL, BLUE, ORANGE = '#42D7B4', '#7FB0ED', '#F0B360'
GRAY, RED = '#8D9BAC', '#E5999F'
FONT = 'DejaVu Sans'
STORY = json.loads((ROOT/'storyboard.json').read_text())
DATA = json.loads((ROOT/'assets/experiment_data.json').read_text())
ROBOT = json.loads((ROOT/'assets/robot_results.json').read_text())
SLOTS = []
BOUNDS = []

for name, weight in [('Deck','normal'),('DeckBold','bold')]:
    pdfmetrics.registerFont(TTFont(name, findfont(FontProperties(family=FONT, weight=weight))))


def rgb(color):
    return RGBColor.from_string(color.lstrip('#'))


def timecode(sec):
    return f'{sec//60:02d}:{sec%60:02d}'


def wrap(text, size, width, bold=False):
    font = 'DeckBold' if bold else 'Deck'
    lines = []
    for para in text.split('\n'):
        if not para:
            lines.append('')
            continue
        line = ''
        for word in para.split():
            candidate = (line+' '+word).strip()
            if line and pdfmetrics.stringWidth(candidate, font, size) > width:
                lines.append(line)
                line = word
            else:
                line = candidate
        lines.append(line)
    return lines


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Pt(W), Pt(H)
        self.prs.core_properties.title = STORY['title']
        self.prs.core_properties.subject = 'Three-minute ICRA supplementary-video storyboard'
        self.prs.core_properties.author = 'CAST'
        self.pdf = pdfcanvas.Canvas(str(ROOT/'cast_video_slides.pdf'), pagesize=(W,H))
        self.pdf.setTitle(STORY['title'])
        self.pdf.setAuthor('CAST')
        self.slide = None
        self.num = 0
        self.prefix = ''
        self.shape_serial = 0

    def name(self, shape, custom=None):
        self.shape_serial += 1
        shape.name = custom or f'{self.prefix}S{self.num:02d}_{self.shape_serial:03d}'
        return shape

    def check(self,x,y,w,h,label):
        assert min(x,y,w,h) >= -0.5, (self.num,label,x,y,w,h)
        assert x+w <= W+0.5 and y+h <= H+0.5, (self.num,label,x,y,w,h)
        BOUNDS.append({'slide':self.num,'type':label,'box':[x,y,w,h]})

    def rect(self,x,y,w,h,fill=CARD,stroke=None,radius=0,name=None,stroke_width=1):
        self.check(x,y,w,h,'shape')
        typ = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        sh=self.name(self.slide.shapes.add_shape(typ,Pt(x),Pt(y),Pt(w),Pt(h)),name)
        if radius:
            sh.adjustments[0]=min(radius/min(w,h),0.2)
        if fill:
            sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
        else:
            sh.fill.background()
        if stroke:
            sh.line.color.rgb=rgb(stroke); sh.line.width=Pt(stroke_width)
        else:
            sh.line.fill.background()
        self.pdf.setFillColor(HexColor(fill or BG))
        self.pdf.setStrokeColor(HexColor(stroke or fill or BG))
        self.pdf.setLineWidth(stroke_width)
        if radius:
            self.pdf.roundRect(x,H-y-h,w,h,radius,fill=int(bool(fill)),stroke=int(bool(stroke)))
        else:
            self.pdf.rect(x,H-y-h,w,h,fill=int(bool(fill)),stroke=int(bool(stroke)))
        return sh

    def ellipse(self,x,y,w,h,fill=None,stroke=None,stroke_width=1):
        self.check(x,y,w,h,'ellipse')
        sh=self.name(self.slide.shapes.add_shape(MSO_SHAPE.OVAL,Pt(x),Pt(y),Pt(w),Pt(h)))
        if fill:
            sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
        else: sh.fill.background()
        if stroke:
            sh.line.color.rgb=rgb(stroke); sh.line.width=Pt(stroke_width)
        else: sh.line.fill.background()
        self.pdf.setFillColor(HexColor(fill or BG)); self.pdf.setStrokeColor(HexColor(stroke or BG))
        self.pdf.setLineWidth(stroke_width)
        self.pdf.ellipse(x,H-y-h,x+w,H-y,fill=int(bool(fill)),stroke=int(bool(stroke)))
        return sh

    def text(self,x,y,w,h,text,size=18,color=WHITE,bold=False,align='left',leading=1.17,name=None):
        self.check(x,y,w,h,'text')
        lines=wrap(text,size,w,bold)
        needed=size*leading*len(lines)
        assert needed <= h+2, f'Slide {self.num} text overflow {needed:.1f}>{h}: {text}'
        font='DeckBold' if bold else 'Deck'
        for line in lines:
            assert pdfmetrics.stringWidth(line,font,size) <= w+0.5,(self.num,line,w)
        box=self.name(self.slide.shapes.add_textbox(Pt(x),Pt(y),Pt(w),Pt(h)),name)
        tf=box.text_frame; tf.clear(); tf.word_wrap=False
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        tf.vertical_anchor=MSO_ANCHOR.TOP
        for i,line in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}[align]
            p.space_before=p.space_after=Pt(0); p.line_spacing=Pt(size*leading)
            r=p.add_run(); r.text=line
            r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=rgb(color)
        self.pdf.setFont(font,size); self.pdf.setFillColor(HexColor(color))
        for i,line in enumerate(lines):
            # Match DejaVu's visible baseline placement in a top-aligned textbox.
            baseline=H-y-size*0.96-i*size*leading
            if align=='center': self.pdf.drawCentredString(x+w/2,baseline,line)
            elif align=='right': self.pdf.drawRightString(x+w,baseline,line)
            else: self.pdf.drawString(x,baseline,line)
        return box

    def line(self,x1,y1,x2,y2,color=LINE,width=1,dash=False):
        sh=self.name(self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Pt(x1),Pt(y1),Pt(x2),Pt(y2)))
        sh.line.color.rgb=rgb(color); sh.line.width=Pt(width)
        if dash:
            from pptx.enum.dml import MSO_LINE_DASH_STYLE
            sh.line.dash_style=MSO_LINE_DASH_STYLE.DASH
        self.pdf.setStrokeColor(HexColor(color)); self.pdf.setLineWidth(width)
        self.pdf.setDash([4,4] if dash else [])
        self.pdf.line(x1,H-y1,x2,H-y2); self.pdf.setDash([])
        return sh

    def arrow(self,x1,y1,x2,y2,color=TEAL,width=2,head=8):
        self.line(x1,y1,x2,y2,color,width)
        ang=math.atan2(y2-y1,x2-x1)
        for sign in [-1,1]:
            back=ang+math.pi+sign*0.45
            self.line(x2,y2,x2+head*math.cos(back),y2+head*math.sin(back),color,width)

    def image(self,path,x,y,w,h,cover=True):
        self.check(x,y,w,h,'image')
        im=Image.open(path).convert('RGB')
        iw,ih=im.size
        if cover:
            ratio=w/h
            if iw/ih>ratio:
                nw=ih*ratio; left=(iw-nw)/2; crop=(left,0,left+nw,ih)
            else:
                nh=iw/ratio; top=(ih-nh)/2; crop=(0,top,iw,top+nh)
            im=im.crop(tuple(round(v) for v in crop))
        else:
            scale=min(w/iw,h/ih); dw,dh=iw*scale,ih*scale
            x+=(w-dw)/2; y+=(h-dh)/2; w,h=dw,dh
        buf=io.BytesIO(); im.save(buf,format='PNG'); buf.seek(0)
        self.name(self.slide.shapes.add_picture(buf,Pt(x),Pt(y),width=Pt(w),height=Pt(h)))
        self.pdf.drawImage(ImageReader(im),x,H-y-h,width=w,height=h)

    def math(self,formula,x,y,w,h,fontsize=24):
        key=hashlib.sha256((formula+str(fontsize)).encode()).hexdigest()[:12]
        path=ROOT/'assets/math'/f'{key}.png'
        if not path.exists():
            fig=plt.figure(figsize=(max(1,w/72),max(0.4,h/72)),dpi=240)
            fig.patch.set_alpha(0)
            fig.text(0.5,0.5,'$'+formula+'$',fontsize=fontsize,color=WHITE,ha='center',va='center')
            fig.savefig(path,transparent=True,bbox_inches='tight',pad_inches=0.02)
            plt.close(fig)
        im=Image.open(path)
        # Flatten on the slide background so PDF and PPTX match reliably.
        bg=Image.new('RGBA',im.size,CARD)
        bg.alpha_composite(im.convert('RGBA'))
        flat=path.with_name(path.stem+'_flat.png'); bg.convert('RGB').save(flat)
        self.image(flat,x,y,w,h,cover=False)

    def begin(self,item):
        self.num=int(item['id']); self.item=item; self.shape_serial=0
        self.slide=self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.slide.background.fill.solid(); self.slide.background.fill.fore_color.rgb=rgb(BG)
        self.pdf.setFillColor(HexColor(BG)); self.pdf.rect(0,0,W,H,fill=1,stroke=0)
        self.pdf.bookmarkPage(item['id']); self.pdf.addOutlineEntry(item['title'],item['id'])
        self.text(42,24,670,18,item['section'].upper(),11,TEAL,bold=True)
        phase='IDEA + METHOD' if item['start']<120 else 'REAL ROBOT'
        self.text(748,24,170,18,phase,10,MUTED,align='right')
        if self.num!=1 and self.num!=10:
            self.text(42,62,876,79,item['title'],31,WHITE,bold=True)

    def embed_completed_media(self):
        for slot in [s for s in SLOTS if s['slide']==self.num]:
            movie=ROOT/slot['filename']
            poster=movie.parent/'posters'/(movie.stem+'.png')
            if not movie.exists() or not poster.exists():
                continue
            x,y,w,h=slot['box_points']
            self.rect(x,y,w,h,CARD,name=f"MEDIA_{slot['id']}_MASK")
            iw,ih=Image.open(poster).size
            scale=min(w/iw,h/ih);vw,vh=iw*scale,ih*scale
            vx,vy=x+(w-vw)/2,y+(h-vh)/2
            self.image(poster,vx,vy,vw,vh,cover=False)
            self.check(vx,vy,vw,vh,'video')
            sh=self.slide.shapes.add_movie(str(movie),Pt(vx),Pt(vy),Pt(vw),Pt(vh),poster_frame_image=str(poster),mime_type='video/mp4')
            self.name(sh,f"MEDIA_{slot['id']}_VIDEO")
            # python-pptx creates click-start video nodes; start at slide entry.
            for cond in self.slide._element.xpath('.//p:video/p:cMediaNode/p:cTn/p:stCondLst/p:cond'):
                cond.set('delay','0')
            slot.update(status='completed; embedded MP4 with poster',poster=str(poster.relative_to(ROOT)),
                        playback_box_points=[vx,vy,vw,vh],pptx_video_shape_name=sh.name)

    def finish(self):
        item=self.item
        self.embed_completed_media()
        self.text(42,512,760,15,'CAST  /  ICRA 2027 supplementary video',9,MUTED)
        self.text(858,512,60,15,f"{self.num:02d} / 10",9,MUTED,align='right')
        self.rect(0,536,W,4,CARD)
        self.rect(0,536,W*(item['start']+item['duration'])/180,4,TEAL)
        slot_notes='\n'.join(f"{s['id']}: {s['brief']}" for s in SLOTS if s['slide']==self.num)
        notes=(f"{timecode(item['start'])}–{timecode(item['start']+item['duration'])} | {item['duration']} seconds\n\n"
               f"NARRATION\n{item['narration']}\n\nVISUAL CUES\n{slot_notes or 'Static slide; no animation required.'}\n\n"
               f"SOURCES\n"+'\n'.join('_ICRA_2027__CAST/'+s for s in item['sources']))
        self.slide.notes_slide.notes_text_frame.text=notes
        trans=OxmlElement('p:transition'); trans.set('advClick','1'); trans.set('advTm',str(item['duration']*1000))
        self.slide._element.insert_element_before(trans,'p:timing','p:extLst')
        self.pdf.showPage()

    def slot(self,ident,x,y,w,h,title,brief,filename,kind='animation'):
        sh=self.rect(x,y,w,h,CARD,LINE,radius=8,name=f'MEDIA_{ident}_FRAME')
        self.text(x+16,y+12,w-32,20,f'{ident}  /  {title}',11,TEAL,bold=True,name=f'MEDIA_{ident}_LABEL')
        self.text(x+16,y+h-23,w-32,16,'STATIC PREVIEW · '+('ANIMATION RESERVED' if kind=='animation' else 'VIDEO RESERVED'),9,MUTED,name=f'MEDIA_{ident}_STATUS')
        SLOTS.append({'id':ident,'slide':self.num,'kind':kind,'title':title,'brief':brief,'filename':filename,
                      'slide_start_seconds':self.item['start'],'slide_duration_seconds':self.item['duration'],
                      'box_points':[x,y,w,h],'box_inches':[round(v/72,5) for v in [x,y,w,h]],
                      'suggested_pixels':[round(w*2),round(h*2)],'pptx_shape_name':sh.name,
                      'status':'reserved; no completed animation or autonomous footage embedded'})
        return x,y,w,h

    def save(self):
        self.prs.save(ROOT/'cast_video_slides.pptx'); self.pdf.save()


def note_card(d,x,y,w,title,body,color=TEAL,h=92):
    d.rect(x,y,w,h,CARD,radius=7)
    d.rect(x,y,4,h,color)
    d.text(x+16,y+12,w-32,27,title,18,color,bold=True)
    d.text(x+16,y+42,w-32,h-47,body,14,MUTED)


def robot_path(task,stage,view):
    return ROOT/'assets/robot'/f'{task}_{stage}_{view}.png'


def point(d,x,y,color,r=5):
    d.ellipse(x-r,y-r,2*r,2*r,color)


def robot_chart(d,task,x,y,w,h):
    values=[100*r['successes']/r['trials'] for r in ROBOT['tasks'][task]['results']]
    colors=[GRAY,ORANGE,BLUE,TEAL]
    d.text(x,y,w,30,'Autonomous success',19,WHITE,bold=True)
    plot_x=x+72; plot_w=w-113
    for idx,(label,val,color) in enumerate(zip(['Base','BC','DICE-RL','CAST'],values,colors)):
        yy=y+54+idx*41
        d.text(x,yy+3,66,24,label,15,color,bold=label=='CAST')
        d.rect(plot_x,yy,plot_w,24,PANEL)
        d.rect(plot_x,yy,plot_w*val/100,24,color)
        d.text(plot_x+plot_w+9,yy+3,42,24,f'{val:.0f}%',15,color,bold=label=='CAST')
    return values


def make_slides(d):
    # 01 — a concise opening; the image is identified as a demonstration.
    d.begin(STORY['slides'][0])
    d.text(42,90,490,103,'CAST',79,TEAL,bold=True)
    d.text(45,201,490,77,'Anchoring critic-guided fine-tuning\nof generative robot policies',24,WHITE)
    d.text(45,322,465,102,'Improve the skill.\nKeep what works.',34,WHITE,bold=True)
    d.rect(554,92,364,344,CARD,radius=10)
    d.image(robot_path('stack','align','chest'),568,107,336,247)
    d.text(570,372,330,45,'OpenArm · staircase manipulation',17,WHITE,bold=True)
    d.text(570,411,330,17,'Teleoperated demonstration frame',10,MUTED)
    d.text(45,462,835,24,'Persistent reference  +  useful adaptation',19,MUTED)
    d.finish()

    # 02 — paired, clearly illustrative concept placeholder.
    d.begin(STORY['slides'][1])
    d.text(44,113,868,27,'An imperfect critic can turn local corrections into cumulative drift.',17,MUTED)
    d.slot('A01',42,153,604,339,'Accumulated drift','Animate matched synthetic actions with clipping alone versus the exact CAST restoration and caps. Same base, critic, and step cap.','animations/A01_accumulated_drift.mp4')
    for pane,heading,color in [(0,'Clipping only',ORANGE),(1,'With restoration',TEAL)]:
        xx=60+pane*296
        d.text(xx+5,199,276,25,heading,18,color,bold=True)
        bx,by=xx+73,348
        d.ellipse(bx-53,by-36,106,72,PANEL,BLUE)
        for dx,dy in [(-21,0),(-11,-10),(0,5),(16,-2),(6,17)]: point(d,bx+dx,by+dy,BLUE,3)
        d.text(xx+22,392,200,23,'Frozen base',14,BLUE)
        if pane==0:
            pts=[(bx,by),(bx+36,by-23),(bx+74,by-49),(bx+112,by-80),(bx+158,by-108)]
        else:
            pts=[(bx,by),(bx+36,by-23),(bx+51,by-37),(bx+45,by-25),(bx+51,by-29)]
        for a,b in zip(pts[:-1],pts[1:]): d.arrow(*a,*b,color=color,width=2,head=5)
        point(d,*pts[-1],color,5)
        if pane==1: d.line(bx,by,*pts[-1],BLUE,1,dash=True)
    d.line(344,199,344,439,LINE,1)
    d.text(71,443,540,22,'Illustrative action space · not a measured learning curve',12,MUTED,align='center')
    note_card(d,669,172,249,'Clipping','Limits each individual correction.',ORANGE,119)
    note_card(d,669,314,249,'Restoration','Responds to departure already accumulated.',TEAL,139)
    d.finish()

    # 03 — editable architecture and same-noise pairing preview.
    d.begin(STORY['slides'][2])
    d.text(43,116,874,24,'Use the same observation and noise to pair the base and current action.',16,MUTED)
    d.rect(43,159,399,304,CARD,radius=8)
    d.text(64,181,357,30,'Observation s + shared noise z',18,WHITE,align='center')
    d.arrow(160,216,160,241,BLUE); d.arrow(335,216,335,241,TEAL)
    d.rect(66,248,167,83,PANEL,BLUE,radius=7)
    d.text(78,264,143,26,'Base policy',19,BLUE,bold=True,align='center')
    d.text(78,295,143,21,'Frozen',15,MUTED,align='center')
    d.rect(253,248,167,83,PANEL,TEAL,radius=7)
    d.text(265,264,143,26,'Residual',19,TEAL,bold=True,align='center')
    d.text(265,295,143,21,'Learned',15,MUTED,align='center')
    d.arrow(150,336,213,367,BLUE); d.arrow(335,336,270,367,TEAL)
    d.math(r'a_\theta(s,z)=\pi_0(s,z)+r_\theta(s,z)',64,376,356,47,fontsize=22)
    d.text(64,433,357,23,'The base supplies samples, not likelihoods.',13,MUTED,align='center')
    d.slot('A02',468,159,450,304,'Paired residual actions','Reveal base actions and learned residual arrows for the same noise samples. Keep the base fixed.','animations/A02_paired_residuals.mp4')
    d.text(489,203,408,25,'Same noise → consistent reference',17,WHITE,bold=True,align='center')
    for bx,by,tx,ty in [(527,290,638,270),(543,343,665,335),(563,385,704,376)]:
        point(d,bx,by,BLUE,5); d.arrow(bx+8,by,tx-8,ty,TEAL,2); point(d,tx,ty,GRAY,5)
    d.text(737,282,151,23,'Base action',15,BLUE)
    d.text(737,318,151,44,'Residual\ncorrection',15,TEAL)
    d.text(737,381,151,25,'Current action',14,GRAY)
    d.finish()

    # 04 — exact target-construction geometry, awaiting animation.
    d.begin(STORY['slides'][3])
    note_card(d,43,154,274,'1  Critic guidance','Normalize and cap the proposal; then scale it.',ORANGE,89)
    note_card(d,43,257,274,'2  Restoration','Weak pull + stronger pull beyond the reference radius.',TEAL,96)
    note_card(d,43,367,274,'3  Combined cap','Limit the final target step from the current action.',BLUE,95)
    d.slot('A03',342,145,576,324,'Construct the action target','Animate the normalized critic proposal, paired weak/radial restoration, and combined displacement cap. Restoration is evaluated at the current action.','animations/A03_cast_update.mp4')
    base=(442,345); current=(584,296)
    d.ellipse(base[0]-65,base[1]-65,130,130,None,BLUE,1)
    d.ellipse(current[0]-60,current[1]-60,120,120,None,MUTED,1)
    point(d,*base,BLUE,6); point(d,*current,GRAY,6)
    d.line(*base,*current,BLUE,1,dash=True)
    proposed=(746,233); restored=(683,255); target=(639,273)
    d.arrow(*current,*proposed,ORANGE,3)
    d.arrow(*proposed,*restored,TEAL,3)
    d.arrow(*current,*target,TEAL,3)
    point(d,*target,TEAL,7)
    d.text(386,415,120,20,'Paired base',14,BLUE)
    d.text(559,370,158,23,'Current action',14,MUTED)
    d.text(702,194,163,25,'Critic proposal',14,ORANGE)
    d.text(743,280,150,43,'Restoring\ncorrection',14,TEAL)
    d.text(649,325,182,25,'Capped target',15,TEAL,bold=True)
    d.text(351,480,558,25,'Targets are constrained; fitted outputs are not hard-bounded.',13,MUTED,align='center')
    d.finish()

    # 05 — fit/refresh and an explanatory penalty curve.
    d.begin(STORY['slides'][4])
    d.slot('A04',42,150,494,292,'Fit → refresh targets','Show fixed targets during regression; move the actor toward them, then recompute. Illustrate the soft restoration potential without claiming a hard policy-output bound.','animations/A04_regress_refresh.mp4')
    d.text(65,194,444,25,'Targets stay fixed during the actor update.',16,WHITE,align='center')
    for xx,yy,tx,ty in [(113,282,189,271),(142,326,201,302),(260,316,287,285),(362,300,345,264),(412,351,375,318)]:
        point(d,xx,yy,GRAY,5); d.arrow(xx+5,yy-2,tx-6,ty+2,TEAL,1.7,6)
        d.rect(tx-5,ty-5,10,10,TEAL)
    d.text(72,372,420,27,'Regress the residual → build fresh targets',16,TEAL,align='center')
    d.rect(559,150,359,292,CARD,radius=8)
    d.text(580,173,319,50,'Restoration penalizes\ndeparture from the base.',20,WHITE,bold=True)
    d.text(582,233,310,22,'Illustrative penalty U(r)',13,MUTED)
    px,py,pw,ph=602,376,278,94
    d.line(px,py,px+pw,py,MUTED,1); d.line(px,py,px,py-ph,MUTED,1)
    rho=0.35
    vals=np.linspace(0,1,36); potential=0.15*vals**2+1.3*np.maximum(vals-rho,0)**2
    points=[(px+pw*v,py-ph*p/potential[-1]) for v,p in zip(vals,potential)]
    for a,b in zip(points[:-1],points[1:]): d.line(*a,*b,TEAL,2.5)
    d.line(px+pw*rho,py+2,px+pw*rho,py-ph,LINE,1,dash=True)
    d.text(px+pw*rho-20,py+5,80,21,'radius',11,MUTED)
    d.text(602,408,280,22,'Soft threshold, not a hard wall',13,TEAL)
    d.text(44,460,872,42,'With inactive caps: the same local gradient direction as regularized critic optimization.',16,MUTED,align='center')
    d.finish()

    # 06 — no fabricated training curves; measured threshold comparison stays visible.
    d.begin(STORY['slides'][5])
    d.slot('A05',42,147,463,303,'Benchmark tasks and measured ablation','First nine seconds: LIBERO and robomimic human demonstration replays, then DMC with a scripted controller. Remaining nineteen seconds: the unchanged measured five-task LIBERO component comparison. Task footage illustrates the environments.','footage/A05_environments_and_results.mp4' if (ROOT/'footage/A05_environments_and_results.mp4').exists() else 'animations/A05_measured_learning_progress.mp4')
    d.text(69,197,409,27,'Success over training',19,WHITE,bold=True)
    d.line(89,379,468,379,LINE,1); d.line(89,246,89,379,LINE,1)
    d.line(89,263,468,263,LINE,1,dash=True)
    d.text(111,278,334,52,'Insert measured evaluation curves',20,MUTED,align='center')
    d.text(113,345,332,24,'Curve animation reserved',13,TEAL,align='center')
    d.text(144,398,279,21,'Environment steps',13,MUTED,align='center')
    d.text(545,153,369,52,'Steps to 90% success',23,WHITE,bold=True)
    runs=DATA['diffusion_efficiency']['environment_steps']
    means={k:np.mean(v)/1000 for k,v in runs.items()}
    px=646; maxw=210
    for method,color,yy in [('CAST',TEAL,233),('DICE-RL',BLUE,293)]:
        d.text(547,yy+4,92,27,method,18,color,bold=True)
        d.rect(px,yy,maxw*means[method]/300,29,color)
        d.text(px+maxw*means[method]/300+9,yy+3,75,27,f'{means[method]:.1f}K',17,WHITE)
    gain=100*(1-means['CAST']/means['DICE-RL'])
    d.text(548,359,145,58,f'{gain:.0f}%',44,TEAL,bold=True)
    d.text(697,369,205,44,'fewer environment\nsteps on average',17,WHITE)
    d.text(547,421,365,22,'Shared diffusion base · three runs',12,MUTED)
    d.text(45,468,870,37,'LIBERO: clipping without restoration loses all success on five tested tasks.',16,WHITE,align='center')
    d.finish()

    # 07 — physical pipeline, with no ambiguous demonstration/evaluation counts.
    d.begin(STORY['slides'][6])
    d.image(robot_path('jigglypuff','initial','chest'),43,151,455,277)
    d.rect(43,390,455,38,CARD)
    d.text(58,398,425,22,'OpenArm + parallel gripper',17,WHITE,bold=True)
    d.image(robot_path('stack','initial','chest'),526,151,183,137)
    d.image(robot_path('stack','grasp','left_wrist'),734,151,183,137)
    d.text(526,299,183,24,'Chest RGB',16,BLUE,align='center')
    d.text(734,299,183,24,'Wrist RGB',16,BLUE,align='center')
    d.rect(526,344,391,91,CARD,radius=7)
    d.text(544,357,355,25,'RGB + joint / gripper state',17,WHITE)
    d.text(544,395,355,24,'Flow policy + residual → action chunks',15,TEAL)
    d.text(46,457,870,30,'Recorded rollouts → offline refinement → autonomous evaluation',20,WHITE,align='center')
    d.text(46,492,870,17,'Scene images are teleoperated demonstration frames.',10,MUTED,align='center')
    d.finish()

    # 08/09 — video space on the left, all four evaluation methods on the right.
    for ident,task,stage,view,slot,title,brief,delta in [
        (8,'jigglypuff','align','chest','V01','Autonomous placement','Insert an actual autonomous placement rollout. Current still is a labeled teleoperated illustration, not a CAST evaluation.',40),
        (9,'stack','align','chest','V02','Autonomous stacking','Insert an actual autonomous stacking rollout showing release and stable final stack. Current still is teleoperated; no per-method clip is available here.',35)]:
        d.begin(STORY['slides'][ident-1])
        d.slot(slot,42,151,542,305,title,brief,f'animations/{slot}_{task}_autonomous.mp4','robot_video')
        d.image(robot_path(task,stage,view),55,189,516,233)
        d.rect(65,198,373,29,BG)
        d.text(75,205,351,19,'PLACEHOLDER · teleoperation still',12,WHITE,bold=True)
        taskkey='jigglypuff_into_case' if task=='jigglypuff' else 'red_on_purple_stacking'
        robot_chart(d,taskkey,610,154,307,231)
        d.text(611,398,298,38,f'+{delta} percentage points',20,TEAL,bold=True)
        d.text(612,436,296,23,'over the pretrained base',15,MUTED)
        endline='Place the toy fully inside the case.' if ident==8 else 'A failed release does not make every earlier action wrong.'
        d.text(44,476,871,28,endline,18,WHITE,align='center')
        d.finish()

    # 10 — one memorable conclusion.
    d.begin(STORY['slides'][9])
    d.text(43,103,871,142,'Improve the skill.\nKeep a persistent reference.',43,WHITE,bold=True)
    for xx,title,sub,col in [(43,'Critic guidance','Propose an improvement',ORANGE),(343,'Soft restoration','Retain a base reference',TEAL),(643,'Bounded target steps','Moderate each correction',BLUE)]:
        d.rect(xx,308,274,109,CARD,radius=8)
        d.text(xx+17,326,240,52,title,21,col,bold=True)
        d.text(xx+17,388,240,22,sub,14,MUTED)
    d.text(44,465,872,26,'cast2027.github.io',20,TEAL)
    d.finish()


def write_documents():
    assert sum(s['duration'] for s in STORY['slides'])==180
    assert sum(s['duration'] for s in STORY['slides'] if s['start']<120)==120
    expected=0
    for s in STORY['slides']:
        assert s['start']==expected
        expected+=s['duration']
    (ROOT/'animation_manifest.json').write_text(json.dumps({'canvas_points':[W,H],'video_canvas_pixels':[1920,1080],'total_seconds':180,'slots':SLOTS},indent=2)+'\n')
    (ROOT/'layout_manifest.json').write_text(json.dumps(BOUNDS,indent=2)+'\n')
    lines=['# CAST — three-minute narration and edit plan','', 'The first 120 seconds explain the idea, method, and simulation evidence. The final 60 seconds cover the real robot and takeaway. Timings are embedded in the PowerPoint. These are draft readings, not a recorded voiceover.','', '| Slide | Time | Duration | Topic | Media |','|---|---|---:|---|---|']
    for s in STORY['slides']:
        lines.append(f"| {s['id']} | {timecode(s['start'])}–{timecode(s['start']+s['duration'])} | {s['duration']} s | {s['title']} | {', '.join(s['slots']) or 'Static'} |")
    lines+=['','## Narration and sources','']
    for s in STORY['slides']:
        words=len(s['narration'].split()); wpm=round(words/s['duration']*60)
        lines += [f"### Slide {s['id']} — {timecode(s['start'])}–{timecode(s['start']+s['duration'])}",'',s['narration'],'',f"Draft pace: {words} words / approximately {wpm} words per minute.",'', 'Sources: '+', '.join('`_ICRA_2027__CAST/'+p+'`' for p in s['sources'])+'.','']
        for slot in SLOTS:
            if slot['slide']==int(s['id']): lines += [f"**{slot['id']} cue:** {slot['brief']}",'']
    lines += ['## Protocol counts awaiting author confirmation','', 'The main physical section says 80 stacking pretraining demonstrations and 30 evaluation attempts. The bundled supplement and result figure retain approximately 60 and 20. This deck deliberately omits those counts. It preserves the reported success percentages (55/80/90/95 for placement; 35/65/60/70 for stacking), rather than inventing new 30-trial counts.','', 'Demonstration stills illustrate the scene only. Replace V01/V02 with correctly attributed autonomous footage before final video export.','']
    (ROOT/'narration.md').write_text('\n'.join(lines))


def render_previews():
    subprocess.run(['pdftoppm','-r','144','-png',str(ROOT/'cast_video_slides.pdf'),str(ROOT/'previews/slide')],check=True)
    paths=sorted((ROOT/'previews').glob('slide-*.png'))
    assert len(paths)==10
    sheet=Image.new('RGB',(1008,1460),BG)
    for i,path in enumerate(paths):
        im=Image.open(path); assert im.size==(1920,1080)
        im.thumbnail((480,270))
        sheet.paste(im,(16+(i%2)*496,16+(i//2)*288))
    sheet.save(ROOT/'previews/contact_sheet.jpg',quality=92)


def main():
    parser=argparse.ArgumentParser(); parser.add_argument('--no-preview',action='store_true')
    args=parser.parse_args()
    for rel in ['assets/math','previews']:(ROOT/rel).mkdir(parents=True,exist_ok=True)
    d=Deck(); make_slides(d); d.save(); write_documents()
    if not args.no_preview: render_previews()
    print('Built 10 editable slides and PDF: 180 seconds (120 method + 60 robot).')
    print(f"Embedded {sum(s['status'].startswith('completed') for s in SLOTS)} animations; remaining media slots stay reserved.")
    print('Narration words:',sum(len(s['narration'].split()) for s in STORY['slides']))


if __name__=='__main__': main()
